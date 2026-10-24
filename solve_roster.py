"""生成月度班表：从三表工作簿 + 月度参数.json 读数据，求解，出报告并导出。

用法:
    python3 solve_roster.py 2026年10月排班表.xlsx 月度参数_10月.json
    python3 solve_roster.py 2026年10月排班表.xlsx 月度参数_10月.json 60
    python3 solve_roster.py 2026年10月排班表.xlsx 月度参数_10月.json 60 anchor
        anchor = 以班表 sheet 的现有内容为参考，尽量少改动

现场调权重（不用改 JSON）:
    python3 solve_roster.py 2026年10月排班表.xlsx 月度参数_10月.json --w R7=15
    python3 solve_roster.py 2026年10月排班表.xlsx 月度参数_10月.json --w R7=15,R4b=0
        每次运行会自动和上一次对比，并把结果另存为带权重标记的文件，多跑几版不会互相覆盖。
        确定好的权重再抄回 JSON 的 weights 里。
"""
import argparse
import json
import os
import sys
import time
from collections import Counter

from roster_rules import (Roster, check_all, score, feasibility_precheck, solve,
                          HARD, SOFT, EARLY, MID, LATE, OFF, CN, GROUP_MGMT,
                          R5_GROUPS, R6_GROUPS, GROUP_SMART)
from roster_io import load_all, RosterDataError


def _w(text):
    """字符串在终端里占几列：中文等全角字符占 2 列，ASCII 占 1 列。
    Python 的 f"{x:>3}" 按字符数补空格，不看显示宽度，中文表格因此会串列。"""
    import unicodedata
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text)


def _pad(text, width, right=True):
    """按显示宽度补足到 width 列。"""
    gap = max(0, width - _w(text))
    return (" " * gap + text) if right else (text + " " * gap)


def export_xlsx(r, human, path, summary=None, requests=None, unmet=None):
    """导出给人事专员用的班表。

    - 格子里写公司统一代码（SFT6 / HTDAYE / CBGDAYK / "/"），可直接贴进公司系统
    - 员工需求格标红：满足的只标红；没满足的标红并加批注写明原因
    - 底部统计行不含管理组（销售口径），在岗总数单列一行为全店口径
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    summary = summary or {}
    requests = requests or {}
    unmet_map = {(n, d1): (want, got, why) for n, d1, want, got, why in (unmet or [])}

    wb = Workbook(); ws = wb.active; ws.title = "班表"
    fills = {EARLY: "FFF2CC", MID: "C6E0B4", LATE: "D9E1F2", OFF: "FFFFFF"}
    base_font = Font(name="Arial")
    bold = Font(name="Arial", bold=True)
    req_font = Font(name="Arial", bold=True, color="C00000")
    center = Alignment(horizontal="center")

    ws.cell(1, 1, "姓名").font = bold
    for d in range(r.cfg.days):
        c = ws.cell(1, 2 + d, d + 1); c.font = bold; c.alignment = center
        ws.cell(2, 2 + d, r.cfg.dow[d]).alignment = center
    for j, code in enumerate([EARLY, MID, LATE, OFF]):
        c = ws.cell(1, 2 + r.cfg.days + j, code); c.font = bold; c.alignment = center

    for i, n in enumerate(r.names):
        row = 3 + i
        ws.cell(row, 1, n).font = base_font
        want_days = requests.get(n, {})
        for d in range(r.cfg.days):
            v = r.grid[n][d]
            c = ws.cell(row, 2 + d, v)                       # 公司统一代码
            c.alignment = center
            c.fill = PatternFill("solid", fgColor=fills[v])
            c.font = req_font if (d + 1) in want_days else base_font
            if (n, d + 1) in unmet_map:
                want, got, why = unmet_map[(n, d + 1)]
                c.comment = Comment(
                    f"需求未满足\n想要 {CN[want]}（{want}），实排 {CN[got]}（{got}）\n"
                    f"原因：" + "；".join(why),
                    "排班系统")
        for j, code in enumerate([EARLY, MID, LATE, OFF]):
            c = ws.cell(row, 2 + r.cfg.days + j, sum(1 for v in r.grid[n] if v == code))
            c.font = base_font; c.alignment = center

    # 底部统计：早/中/晚 不含管理组；在岗 单列全店口径
    sales = [n for n in r.names if r.group_of[n] != GROUP_MGMT]
    base = 3 + len(r.names)
    stat_rows = [(f"{EARLY}（不含管理组）", lambda d: r.count(d, EARLY, sales)),
                 (f"{MID}（不含管理组）",   lambda d: r.count(d, MID, sales)),
                 (f"{LATE}（不含管理组）",  lambda d: r.count(d, LATE, sales)),
                 ("在岗（全店）",           lambda d: r.on_duty(d))]
    for j, (lab, fn) in enumerate(stat_rows):
        ws.cell(base + j, 1, lab).font = bold
        for d in range(r.cfg.days):
            ws.cell(base + j, 2 + d, fn(d)).alignment = center

    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 20
    note = base + len(stat_rows) + 1
    ws.cell(note, 1, "红字 = 员工需求的班次；带批注的红字 = 需求未满足，鼠标悬停看原因。"
                     "底色：黄=SFT6 绿=HTDAYE 蓝=CBGDAYK 白=休")

    # 记录这一版是用什么参数跑出来的——现场多跑几版时靠它区分
    ps = wb.create_sheet("参数")
    ps.cell(1, 1, "项目").font = bold
    ps.cell(1, 2, "取值").font = bold
    quotas = list(r.cfg.off_quota.values())
    common = max(set(quotas), key=quotas.count) if quotas else 0
    special = "、".join(f"{n}{q}天" for n, q in r.cfg.off_quota.items() if q != common)
    total_req = sum(len(v) for v in requests.values())
    rows = [("生成时间", time.strftime("%Y-%m-%d %H:%M")),
            ("软约束罚分", summary.get("soft", "")),
            ("全员上班日", "、".join(map(str, r.cfg.all_hands_days)) or "无"),
            ("休息额度", f"一般 {common} 天" + (f"；特殊：{special}" if special else "")),
            ("员工需求", f"共 {total_req} 条，满足 {total_req - len(unmet_map)} 条"),
            ("", "")]
    rows += [(f"权重 {k}", v) for k, v in sorted(r.cfg.weights.items())]
    rows += [("", ""), ("违规次数", "")]
    rows += [(f"  {k}", v) for k, v in sorted(summary.get("rules", {}).items())]
    for i, (k, v) in enumerate(rows):
        ps.cell(2 + i, 1, k)
        ps.cell(2 + i, 2, v)
    ps.column_dimensions["A"].width = 16
    ps.column_dimensions["B"].width = 44
    wb.save(path)


def report_unmet(sol, requests):
    """未满足的员工需求 + 归因：如果硬把他调过去，会撞上哪条规则。

    看两头：
      一头是【离开】原来的班次——他走了，原来那个班会不会缺人（R5 / R6 / 周末在岗）
      一头是【去到】想要的班次——能不能接上前后两天（R2 晚接早）、当天能不能休（全员日）
    """
    cfg = sol.cfg
    r5 = set(sol.members(*R5_GROUPS))
    out = []
    for n, days in (requests or {}).items():
        cfg_group = sol.group_of.get(n)
        for d1, want in days.items():
            got = sol.grid[n][d1 - 1]
            if got == want:
                continue
            d, why = d1 - 1, []

            # ---- 去到：想要的班次本身能不能排 ----
            if want == OFF and d1 in cfg.all_hands_days:
                why.append("当日为全员上班日，不能排休")
            if want == EARLY and d > 0 and sol.grid[n][d - 1] == LATE:
                why.append("前一天是晚班，晚接早不允许")
            if want == LATE and d < cfg.days - 1 and sol.grid[n][d + 1] == EARLY:
                why.append("后一天是早班，晚接早不允许")
            if cfg_group == GROUP_SMART and want in (EARLY, MID, LATE):
                c = sol.count(d, want, sol.members(GROUP_SMART))
                if c >= cfg.max_smart_same_shift:
                    why.append(f"智慧办公当日{CN[want]}班已有 {c} 人，达到上限 {cfg.max_smart_same_shift} 人")
            if want == MID:
                mids = sum(1 for v in sol.grid[n] if v == MID)
                if mids >= cfg.max_mid_per_person:
                    why.append(f"本月已有 {mids} 个中班，达到每人上限 {cfg.max_mid_per_person} 个")

            # ---- 离开：原来的班次会不会缺人 ----
            if got in (EARLY, LATE):
                lo = cfg.min_early if got == EARLY else cfg.min_late
                if n in r5:
                    c = sol.count(d, got, list(r5))
                    if c <= lo:
                        why.append(f"当日销售{CN[got]}班仅 {c} 人，调走就低于下限 {lo}")
                if cfg_group in R6_GROUPS:
                    who = sol.members(cfg_group)
                    if sol.count(d, got, who) <= 1:
                        why.append(f"此人是{cfg_group}当日唯一的{CN[got]}班")
            if want == OFF and got != OFF and sol.is_weekend(d):
                on = sol.on_duty(d)
                if on <= cfg.weekend_min_on:
                    why.append(f"周末在岗仅 {on} 人，已达下限 {cfg.weekend_min_on}")
            if got == OFF and want != OFF:
                why.append(f"这天被安排为休息（每人休息需排满 {cfg.off_quota[n]} 天）")

            out.append((n, d1, want, got, why or ["与其他规则综合权衡后未能满足"]))
    return out


SNAP = ".上次运行.json"      # 隐藏文件，存上一次的结果用来对比


def parse_overrides(text):
    """'R7=15,R4b=0' -> {'R7': 15, 'R4b': 0}"""
    out = {}
    for part in (text or "").replace("，", ",").split(","):
        part = part.strip()
        if not part:
            continue
        if "=" not in part:
            raise ValueError(f"权重写法应为 键=数字，收到「{part}」")
        k, v = part.split("=", 1)
        out[k.strip()] = int(v.strip())
    return out


def compare_with_last(month_key, summary, grid):
    """和上一次（同一个月）的运行结果对比。"""
    last = None
    if os.path.exists(SNAP):
        try:
            data = json.load(open(SNAP, encoding="utf-8"))
            if data.get("month") == month_key:
                last = data
        except (OSError, ValueError):
            last = None

    if last:
        print("\n-- 与上一次对比 --")
        changed_w = {k: (last["weights"].get(k), v) for k, v in summary["weights"].items()
                     if last["weights"].get(k) != v}
        if changed_w:
            print("  权重变化: " + "，".join(f"{k} {a}->{b}" for k, (a, b) in changed_w.items()))
        else:
            print("  权重未变")
        d = summary["soft"] - last["soft"]
        print(f"  罚分: {last['soft']} -> {summary['soft']}  ({d:+d})")
        keys = sorted(set(last["rules"]) | set(summary["rules"]))
        moved = [(k, last["rules"].get(k, 0), summary["rules"].get(k, 0)) for k in keys
                 if last["rules"].get(k, 0) != summary["rules"].get(k, 0)]
        for k, a, b in moved:
            print(f"    {k:9s} {a:3d} -> {b:3d}  ({b - a:+d})")
        if not moved:
            print("    各规则违规次数均无变化")
        cells = sum(1 for n, row in grid.items() if n in last["grid"]
                    for a, b in zip(row, last["grid"][n]) if a != b)
        tot = sum(len(row) for row in grid.values())
        print(f"  班表改动: {cells}/{tot} 格")

    snap = dict(summary, month=month_key, grid=grid)
    with open(SNAP, "w", encoding="utf-8") as f:
        json.dump(snap, f, ensure_ascii=False)


def main():
    ap = argparse.ArgumentParser(description="生成月度班表")
    ap.add_argument("xlsx")
    ap.add_argument("params", nargs="?", default="月度参数.json")
    ap.add_argument("secs", nargs="?", type=int, default=60)
    ap.add_argument("mode", nargs="?", default="free", choices=["free", "anchor"])
    ap.add_argument("--w", default="", help="临时覆盖权重，如 R7=15,R4b=0")
    a = ap.parse_args()
    xlsx, params, secs, mode = a.xlsx, a.params, a.secs, a.mode

    try:
        overrides = parse_overrides(a.w)
        cfg, names, gof, req, human = load_all(xlsx, params)
    except (RosterDataError, ValueError) as e:
        print(f"\n[数据错误] {e}")
        sys.exit(1)

    if overrides:
        unknown = [k for k in overrides if k not in cfg.weights]
        if unknown:
            print(f"\n[数据错误] 没有这些权重: {'、'.join(unknown)}"
                  f"\n  可调的有: {'、'.join(sorted(cfg.weights))}")
            sys.exit(1)
        cfg.weights.update(overrides)
        print("\n[现场调参] " + "，".join(f"{k}={v}" for k, v in overrides.items())
              + "（仅本次生效，JSON 未改动）")

    if human:
        hh, hs = score(check_all(human))
        print(f"\n== 原班表基准 ==  硬约束 {hh} 项 / 软约束罚分 {hs}")

    print("\n-- 可行性预检 --")
    for msg in feasibility_precheck(names, gof, cfg) or ["全部通过"]:
        print("  " + msg)

    solver_req = dict(req)
    if mode == "anchor":
        if not human:
            print("\n[anchor 模式需要已有班表，改用 free]"); mode = "free"
        else:
            for n in names:
                solver_req.setdefault(n, {})
                for d in range(cfg.days):
                    solver_req[n].setdefault(d + 1, human.grid[n][d])
            print("\n[anchor 模式] 以原班表为参考，优先少改动")

    print(f"\n-- 求解中（上限 {secs} 秒）--")
    t0 = time.time()
    sol, status = solve(names, gof, cfg, requests=solver_req, max_seconds=secs)
    print(f"  状态: {status}   耗时: {time.time() - t0:.1f}s")

    if sol is None:
        print("\n  无解。建议放宽顺序：")
        print("    1) 减少人事的锁定/禁止格（上方若列出冲突项，优先处理）")
        print("    2) mgmt_avoid_weekend_off 设为 false")
        print("    3) 下调 weekend_min_on / min_early / min_late")
        print("    4) 上调 balance_tol")
        return

    vs = check_all(sol)
    sh, ss = score(vs)
    print(f"\n== 求解版 ==  硬约束 {sh} 项 / 软约束罚分 {ss}")
    by = Counter(v.rule for v in vs)
    if by:
        print("  " + "  ".join(f"{k}x{v}" for k, v in sorted(by.items())))
    for v in [v for v in vs if v.level == HARD]:
        print("  [硬] " + str(v))

    unmet = report_unmet(sol, req) if req else []
    if req:
        okc = sum(len(v) for v in req.values()) - len(unmet)
        print(f"\n-- 员工需求 --  {okc}/{sum(len(v) for v in req.values())} 条已满足")
        for n, d1, want, got, why in unmet:
            print(f"  {n} {d1}日 想要「{CN[want]}」实排「{CN[got]}」")
            for w in why:
                print(f"      原因: {w}")

    mid = sum(1 for n in names for v in sol.grid[n] if v == MID)
    print(f"\n-- 中班用量 --  {mid} 人日（日均 {mid/cfg.days:.1f}，免罚额度 {cfg.mid_soft_cap}/天）")

    if human:
        diff = sum(1 for n in names for d in range(cfg.days)
                   if human.grid[n][d] != sol.grid[n][d])
        tot = len(names) * cfg.days
        print(f"\n-- 改动量 --  {diff}/{tot} 格 ({diff/tot*100:.0f}%)")

    print("\n-- 求解版班表 --")
    NW = max(_w(n) for n in names) + 1            # 姓名列宽，按显示宽度算
    print(" " * NW + "".join(_pad(str(d + 1), 3) for d in range(cfg.days)))
    print(" " * NW + "".join(_pad(cfg.dow[d], 3) for d in range(cfg.days)))
    for n in names:
        print(_pad(n, NW, right=False)
              + "".join(_pad(CN[v], 3) for v in sol.grid[n]))

    summary = {"soft": ss, "hard": sh, "rules": dict(Counter(v.rule for v in vs)),
               "weights": dict(cfg.weights)}
    compare_with_last(f"{cfg.year}-{cfg.month}", summary, sol.grid)

    tag = "_" + "_".join(f"{k}-{v}" for k, v in overrides.items()) if overrides else ""
    out = f"{cfg.year}年{cfg.month}月_求解结果{tag}.xlsx"
    export_xlsx(sol, human, out, summary, requests=req, unmet=unmet)
    print(f"\n已导出: {out}（红字 = 员工需求；带批注 = 需求未满足）")


if __name__ == "__main__":
    main()
